#!/usr/bin/env python3
"""
Exports presentation slides from workspace/index.html to a native PowerPoint (.pptx) file.
Preserves 16:9 widescreen layout, dark theme, slide titles, cards, stats, and Persian RTL text.
"""

import sys
import re
from pathlib import Path
from bs4 import BeautifulSoup
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Theme Colors
BG_DARK = RGBColor(14, 17, 23)
CARD_BG = RGBColor(24, 29, 39)
CARD_BORDER = RGBColor(45, 55, 72)
TEXT_WHITE = RGBColor(248, 250, 252)
TEXT_MUTED = RGBColor(156, 163, 175)
TEXT_DIM = RGBColor(100, 116, 139)

TONE_COLORS = {
    "cyan": RGBColor(56, 189, 248),
    "amber": RGBColor(251, 191, 36),
    "emerald": RGBColor(52, 211, 153),
    "violet": RGBColor(167, 139, 250),
    "rose": RGBColor(251, 113, 133),
}

DEFAULT_FONT = "Vazirmatn"
FALLBACK_FONTS = "Tahoma, Segoe UI, Arial"

def clean_text(text):
    if not text:
        return ""
    return re.sub(r'\s+', ' ', text).strip()

def is_persian_text(text):
    return any(u'\u0600' <= c <= u'\u06FF' for c in text)

def set_shape_bg(shape, rgb_color, border_color=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()

def create_pptx_from_html(html_path, output_path):
    with open(html_path, "r", encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    is_rtl = soup.html and soup.html.get("dir", "").lower() == "rtl"
    default_align = PP_ALIGN.RIGHT if is_rtl else PP_ALIGN.LEFT

    slides = soup.find_all("section", class_="slide")
    if not slides:
        # Fallback if no .slide sections
        slides = soup.find_all("section")

    for idx, slide_elm in enumerate(slides):
        slide = prs.slides.add_slide(blank_layout)

        # Background fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        set_shape_bg(bg_shape, BG_DARK)

        # Tone accent
        tone = slide_elm.get("data-tone", "cyan")
        accent_color = TONE_COLORS.get(tone, TONE_COLORS["cyan"])

        # Top subtle accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Pt(3))
        set_shape_bg(bar, accent_color)

        # Slide Number / Section tracker
        sec_name = slide_elm.get("data-section", "")
        if not sec_name:
            sec_tag = slide_elm.find(class_=re.compile(r'(kicker|badge|section-tag|category|tag)'))
            if sec_tag:
                sec_name = clean_text(sec_tag.get_text())

        # Header Box (Kicker + Main Title + Subtitle)
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(1.8))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Kicker / Badge
        p_kicker = tf.paragraphs[0]
        p_kicker.text = f"{idx + 1:02d}  ·  {sec_name.upper()}" if sec_name else f"SLIDE {idx + 1:02d}"
        p_kicker.font.name = DEFAULT_FONT
        p_kicker.font.size = Pt(11)
        p_kicker.font.bold = True
        p_kicker.font.color.rgb = accent_color
        p_kicker.alignment = default_align

        # Main Title (h1 or h2)
        h_elm = slide_elm.find(["h1", "h2"])
        title_text = clean_text(h_elm.get_text()) if h_elm else f"Slide {idx + 1}"
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.name = DEFAULT_FONT
        p_title.font.size = Pt(26 if len(title_text) > 40 else 32)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.alignment = default_align

        # Optional Subtitle / Lead Paragraph
        sub_elm = slide_elm.find(["p", "h3"], class_=re.compile(r'(lead|subtitle|desc|summary)'))
        if not sub_elm and h_elm:
            next_p = h_elm.find_next_sibling("p")
            if next_p and not next_p.find_parent(class_=re.compile(r'(card|box|stat|item)')):
                sub_elm = next_p

        if sub_elm:
            sub_text = clean_text(sub_elm.get_text())
            if sub_text and sub_text != title_text and len(sub_text) < 180:
                p_sub = tf.add_paragraph()
                p_sub.text = sub_text
                p_sub.font.name = DEFAULT_FONT
                p_sub.font.size = Pt(14)
                p_sub.font.color.rgb = TEXT_MUTED
                p_sub.alignment = default_align

        # Main Content Area: Cards or Grid or Bullet items
        cards = slide_elm.find_all(class_=re.compile(r'(card|panel|stat|item|tile|block)'))
        # Filter out deeply nested child cards
        top_cards = []
        for c in cards:
            if not any(parent in cards for parent in c.parents):
                top_cards.append(c)

        content_top = Inches(2.7)
        content_height = Inches(4.2)
        content_width = Inches(11.733)
        margin_left = Inches(0.8)

        if top_cards:
            count = min(len(top_cards), 4)  # fit up to 4 columns nicely
            gap = Inches(0.25)
            card_w = (content_width - (gap * (count - 1))) / count

            for c_idx in range(count):
                c_elm = top_cards[c_idx]
                c_left = margin_left + c_idx * (card_w + gap)

                # Card shape
                c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, content_top, card_w, content_height)
                set_shape_bg(c_shape, CARD_BG, CARD_BORDER)

                c_tf = c_shape.text_frame
                c_tf.word_wrap = True
                c_tf.vertical_anchor = MSO_ANCHOR.TOP
                c_tf.margin_left = c_tf.margin_right = Inches(0.25)
                c_tf.margin_top = c_tf.margin_bottom = Inches(0.25)

                # Check for big stat / number
                stat_val = c_elm.find(class_=re.compile(r'(number|value|stat|metric|pct|percent)'))
                stat_text = clean_text(stat_val.get_text()) if stat_val else ""
                
                # Card Title / Head
                c_head = c_elm.find(["h3", "h4", "h5", "strong"])
                c_title = clean_text(c_head.get_text()) if c_head else ""

                first_para = True
                if stat_text and len(stat_text) <= 12 and (any(ch.isdigit() for ch in stat_text) or '%' in stat_text):
                    p_stat = c_tf.paragraphs[0]
                    p_stat.text = stat_text
                    p_stat.font.name = DEFAULT_FONT
                    p_stat.font.size = Pt(36)
                    p_stat.font.bold = True
                    p_stat.font.color.rgb = accent_color
                    p_stat.alignment = default_align
                    first_para = False

                if c_title and c_title != stat_text:
                    p_ct = c_tf.paragraphs[0] if first_para else c_tf.add_paragraph()
                    p_ct.text = c_title
                    p_ct.font.name = DEFAULT_FONT
                    p_ct.font.size = Pt(17)
                    p_ct.font.bold = True
                    p_ct.font.color.rgb = TEXT_WHITE
                    p_ct.alignment = default_align
                    first_para = False

                # Card body / list / description
                c_paragraphs = c_elm.find_all(["p", "li"])
                added_lines = 0
                for cp in c_paragraphs:
                    cp_text = clean_text(cp.get_text())
                    if not cp_text or cp_text == c_title or cp_text == stat_text:
                        continue
                    if added_lines >= 5:
                        break
                    p_desc = c_tf.paragraphs[0] if first_para else c_tf.add_paragraph()
                    is_bullet = (cp.name == "li")
                    p_desc.text = f"•  {cp_text}" if is_bullet else cp_text
                    p_desc.font.name = DEFAULT_FONT
                    p_desc.font.size = Pt(13)
                    p_desc.font.color.rgb = TEXT_MUTED
                    p_desc.alignment = default_align
                    first_para = False
                    added_lines += 1

        else:
            # Fallback layout: Bullet points or full-width text blocks
            body_box = slide.shapes.add_textbox(margin_left, content_top, content_width, content_height)
            b_tf = body_box.text_frame
            b_tf.word_wrap = True
            b_tf.margin_left = b_tf.margin_top = b_tf.margin_right = b_tf.margin_bottom = 0

            items = slide_elm.find_all(["li", "p"])
            first_b = True
            added = 0
            for item in items:
                it_text = clean_text(item.get_text())
                if not it_text or it_text == title_text:
                    continue
                if added >= 6:
                    break
                p_it = b_tf.paragraphs[0] if first_b else b_tf.add_paragraph()
                p_it.text = f"•   {it_text}" if item.name == "li" else it_text
                p_it.font.name = DEFAULT_FONT
                p_it.font.size = Pt(16)
                p_it.font.color.rgb = TEXT_WHITE
                p_it.alignment = default_align
                first_b = False
                added += 1

    prs.save(output_path)
    print(f"Exported {len(slides)} slides to {output_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: python3 export_pptx.py <input_html> <output_pptx>")
        sys.exit(1)
    create_pptx_from_html(sys.argv[1], sys.argv[2])
